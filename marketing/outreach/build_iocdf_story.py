from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    "/Users/batman/gentleloop-labs/patterns/outputs/"
    "patterns-outreach-2026-09-08/"
    "IOCDF-Aftaab-Siddiqui-When-Tracking-Becomes-Checking.docx"
)


TITLE = "When Tracking Becomes Checking"


BODY = [
    "My psychologist once asked me to keep track of my obsessions, compulsions, "
    "and distress. The request made sense. By the time I reached an appointment, "
    "the week had blurred. I could remember that Tuesday was difficult, but not "
    "what triggered the distress, how I responded, or how long the relief lasted.",

    "Writing things down gave those moments a shape. I began to see a sequence: a "
    "doubt arrived, distress made it feel urgent, a compulsion promised relief, and "
    "the same doubt eventually returned. A few notes helped me describe that loop "
    "more clearly in therapy. They also helped me notice mental rituals that had "
    "previously looked like ordinary thinking.",

    "Then the record itself started changing. I would write an entry and reread it. "
    "I would reconsider the category, adjust a distress rating, and return later to "
    "make sure I had described the experience correctly. I was no longer preserving "
    "information for a conversation. I was asking the journal to give me the feeling "
    "that nothing important had been missed.",

    "That feeling never held. A more precise entry created another detail to inspect. "
    "A number that felt accurate in the morning looked questionable by evening. The "
    "better I became at documenting the experience, the more material I had available "
    "for review. Something useful had quietly become another checking surface.",

    "The difficult part was that the two actions looked identical from the outside. "
    "Both involved opening a journal and recording an OCD event. The difference was "
    "the job I was giving the action. Sometimes I wanted a brief record that could "
    "help me reflect later. At other times I wanted certainty that the experience had "
    "been captured perfectly. No button could tell those intentions apart for me.",

    "I started asking myself two questions before recording anything: What practical "
    "job does this note have, and how will I know when that job is finished? If I "
    "could not answer, I tried to leave the moment unrecorded. If the note had a clear "
    "purpose, I wrote enough for that purpose and stopped. When the urge to revise it "
    "returned, the unfinished feeling was sometimes the thing I needed to allow.",

    "These were personal boundaries, not a treatment plan for anyone else. They did "
    "not make every decision obvious. I still caught myself checking. The questions "
    "simply gave me a way to notice when I had crossed from remembering into trying "
    "to settle uncertainty.",

    "This experience followed me into software development. I built Patterns because "
    "I wanted a private place to journal and support ERP practice between appointments. "
    "I also had to accept that an OCD tool could participate in the same loop it was "
    "meant to help someone observe.",

    "That changed the design decisions I cared about. Personal records stay on the "
    "device, so an intimate account of intrusive thoughts does not need to become a "
    "cloud profile. An account is not required. Reminders are optional. Entries do not "
    "need to be polished, and missing a day does not erase earlier work. The app does "
    "not contain a chatbot that interprets a thought or supplies reassurance about "
    "what it means.",

    "I also began questioning the ordinary definition of engagement. Product teams "
    "often celebrate more sessions, more taps, and more time spent inside an app. For "
    "this kind of tool, a good session may last less than a minute. Someone records "
    "what matters, chooses a response, and returns to the rest of the day. Keeping "
    "that session going would serve the metric more than the person.",

    "Lived experience helped me recognize this risk, but it did not make me a clinician "
    "or make my preferences universal. Some people find detailed tracking useful. "
    "Others may need less measurement, or a different tool entirely. A responsible "
    "product should leave room for those differences and state its limits clearly.",

    "I still write things down. The practice remains useful when I give it a specific "
    "purpose and an ending. I record enough to remember the pattern, then try to let "
    "the record remain imperfect. The discomfort that follows is not evidence that I "
    "failed to document the moment. It is often the uncertainty I was trying to escape.",

    "The lesson I carry forward is modest: a helpful action can change its function. "
    "The question is not only whether journaling or tracking is good. It is what I am "
    "asking it to do right now, and whether I can stop when that job is done.",
]


BIO = (
    "Aftaab Siddiqui is an independent developer living with OCD and the creator of "
    "Patterns, a privacy-first OCD journal and ERP practice companion. He writes about "
    "living with OCD and building mental-health technology that respects privacy, "
    "uncertainty, and people’s ability to step away."
)


def set_font(run, name="Arial", size=11, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def remove_paragraph_borders(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_font(run, size=9)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.12

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(18)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_borders(title_style)

    heading_style = styles["Heading 1"]
    heading_style.font.name = "Arial"
    heading_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading_style.font.size = Pt(12)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run(TITLE)
    set_font(title_run, size=18, bold=True)
    remove_paragraph_borders(title)
    title.paragraph_format.space_after = Pt(4)

    byline = doc.add_paragraph()
    byline.paragraph_format.space_after = Pt(14)
    run = byline.add_run("By Aftaab Siddiqui")
    set_font(run, size=10, italic=True)

    for text in BODY:
        paragraph = doc.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.widow_control = True

    heading = doc.add_paragraph("Author bio", style="Heading 1")
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(4)

    bio = doc.add_paragraph(BIO)
    bio.paragraph_format.space_after = Pt(0)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    core_props = doc.core_properties
    core_props.title = TITLE
    core_props.author = "Aftaab Siddiqui"
    core_props.subject = "Personal story submission to the International OCD Foundation"
    core_props.keywords = "OCD, personal story, journaling, tracking, uncertainty"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
